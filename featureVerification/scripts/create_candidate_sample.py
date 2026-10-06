import argparse
import json
import random
import statistics
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = REPOSITORY_ROOT / 'verified-features.xlsx'
DEFAULT_OUTPUT = REPOSITORY_ROOT / 'candidate-case-sample.xlsx'
DEFAULT_SAMPLE_SIZE = 150
DEFAULT_SEED = 20261005
ROLE_TARGETS = {
	'Courier': 55,
	'Actual trafficker': 28,
	'Storekeeper': 22,
	'Packager': 18,
	'Manager/organiser': 12,
	'Operator/financial controller': 8,
}
ROLE_MAP = {
	'courier': 'Courier',
	'storekeeper': 'Storekeeper',
	'packager': 'Packager',
	'actual trafficker': 'Actual trafficker',
	'manager/organizer': 'Manager/organiser',
	'manager/organiser': 'Manager/organiser',
	'operator/financial controller': 'Operator/financial controller',
	'international operator/financial controller': 'Operator/financial controller',
}
GROUP_WEIGHTS = {
	'age': 1.2,
	'role': 1.8,
	'district': 1.0,
	'nature': 1.1,
	'subdistrict': 0.65,
	'drug': 1.5,
	'quantity': 1.4,
}


def read_rows(sheet):
	iterator = sheet.iter_rows(values_only=True)
	headers = next(iterator)
	return headers, [dict(zip(headers, row)) for row in iterator]


def parse_array(value):
	if not value:
		return []
	if isinstance(value, list):
		return value
	try:
		result = json.loads(value)
		return result if isinstance(result, list) else []
	except (TypeError, json.JSONDecodeError):
		return []


def normalized_id(value):
	if value is None:
		return None
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value)


def age_band(value):
	try:
		age = float(value)
	except (TypeError, ValueError):
		return None
	if age < 21:
		return 'Under 21'
	if age < 35:
		return '21–34'
	return '35+'


def normalized_role(value):
	role = str(value or '').strip()
	return ROLE_MAP.get(role.casefold(), role)


def build_features(sheets):
	judgments = sheets['judgments'][1]
	charges = sheets['charges'][1]
	charge_defendants = sheets['charge_defendants'][1]
	defendants = sheets['defendants'][1]
	trials = sheets['trials'][1]

	case_rows = {
		row.get('source_judgement_id'): row
		for row in judgments
		if row.get('source_judgement_id') and row.get('exclude') is not True
	}
	eligible = set(case_rows)
	features = {case_id: set() for case_id in eligible}
	profiles = defaultdict(dict)

	for row in defendants:
		case_id = row.get('source_judgement_id')
		if case_id not in eligible:
			continue
		defendant_id = normalized_id(row.get('defendant_id'))
		if defendant_id is not None:
			profiles[case_id][defendant_id] = row.get('age_at_offence.age')
			band = age_band(row.get('age_at_offence.age'))
			if band:
				features[case_id].add(f'age:{band}')

	for row in charges:
		case_id = row.get('source_judgement_id')
		if case_id not in eligible:
			continue
		for column, group in (
			('place_of_offence.nature', 'nature'),
			('place_of_offence.subDistrict', 'subdistrict'),
			('place_of_offence.district', 'district'),
		):
			if row.get(column):
				features[case_id].add(f'{group}:{row[column]}')

	for row in charge_defendants:
		case_id = row.get('source_judgement_id')
		if case_id not in eligible:
			continue
		for item in parse_array(row.get('roles_facts')):
			role = normalized_role(item.get('role'))
			if role in ROLE_TARGETS:
				features[case_id].add(f'role:{role}')

	drug_rows = []
	quantities_by_drug = defaultdict(list)
	for row in trials:
		case_id = row.get('source_judgement_id')
		if case_id not in eligible:
			continue
		for item in parse_array(row.get('drugs')):
			drug = item.get('other_drug_type') if item.get('drug_type') == 'Other' else item.get('drug_type')
			if not drug:
				continue
			drug = str(drug)
			drug_rows.append((row, item, drug))
			features[case_id].add(f'drug:{drug}')
			quantity = item.get('quantity')
			if isinstance(quantity, (int, float)) and quantity > 0:
				quantities_by_drug[drug].append(float(quantity))

	medians = {
		drug: statistics.median(values)
		for drug, values in quantities_by_drug.items()
	}
	for row, item, drug in drug_rows:
		quantity = item.get('quantity')
		if isinstance(quantity, (int, float)) and quantity > 0 and drug in medians:
			band = 'smaller' if quantity <= medians[drug] else 'larger'
			features[row['source_judgement_id']].add(f'quantity:{drug}:{band}')

	return case_rows, eligible, features


def build_targets(features):
	counts = Counter(feature for feature_set in features.values() for feature in feature_set)
	groups = defaultdict(list)
	for feature, count in counts.items():
		group, value = feature.split(':', 1)
		groups[group].append((value, count))
	for group in groups:
		groups[group].sort(key=lambda item: (-item[1], item[0]))

	targets = {}
	for group, values in groups.items():
		if group == 'age':
			targets[group] = {value: min(35, count) for value, count in values}
		elif group == 'role':
			targets[group] = {
				value: min(ROLE_TARGETS[value], count)
				for value, count in values
				if value in ROLE_TARGETS
			}
		elif group == 'district':
			targets[group] = {value: min(4, count) for value, count in values}
		elif group == 'nature':
			targets[group] = {
				value: min(5 if count >= 10 else 2, count)
				for value, count in values
			}
		elif group == 'subdistrict':
			targets[group] = {value: min(2, count) for value, count in values[:30]}
		elif group == 'drug':
			targets[group] = {
				value: min(5 if count >= 100 else 3 if count >= 20 else 1, count)
				for value, count in values
			}
		elif group == 'quantity':
			drug_rank = {
				value: index
				for index, (value, _) in enumerate(groups['drug'])
			}
			targets[group] = {
				value: min(
					8 if drug_rank[value.rsplit(':', 1)[0]] < 6
					else 4 if drug_rank[value.rsplit(':', 1)[0]] < 12
					else 1,
					count,
				)
				for value, count in values
			}
	return targets


def select_cases(features, sample_size, seed):
	targets = build_targets(features)
	randomizer = random.Random(seed)
	remaining = sorted(features)
	selected = []
	selected_counts = Counter()
	while remaining and len(selected) < sample_size:
		best_case = None
		best_score = -1.0
		for case_id in remaining:
			group_scores = []
			for group, group_targets in targets.items():
				values = sorted(
					feature[len(group) + 1:]
					for feature in features[case_id]
					if feature.startswith(group + ':')
					and feature[len(group) + 1:] in group_targets
				)
				if values:
					deficits = [
						max(
							0.0,
							(group_targets[value] - selected_counts[group + ':' + value])
							/ group_targets[value],
						)
						for value in values
						if group_targets[value]
					]
					if deficits:
						group_scores.append(
							GROUP_WEIGHTS[group] * sum(deficits) / len(deficits)
						)
			score = sum(group_scores) + randomizer.random() * 0.05
			if score > best_score:
				best_case, best_score = case_id, score
		selected.append(best_case)
		remaining.remove(best_case)
		selected_counts.update(features[best_case])
	return selected


def filter_workbook(workbook, selected, sample_size):
	selected_set = set(selected)
	original_headers = {}
	row_count = {}
	for sheet in workbook.worksheets:
		headers = [cell.value for cell in sheet[1]]
		original_headers[sheet.title] = headers
		if 'source_judgement_id' not in headers:
			raise ValueError(f'{sheet.title} has no source_judgement_id column')
		id_column = headers.index('source_judgement_id') + 1
		kept_rows = []
		for row_index in range(2, sheet.max_row + 1):
			case_id = sheet.cell(row_index, id_column).value
			if case_id not in selected_set:
				continue
			cells = [sheet.cell(row_index, column) for column in range(1, len(headers) + 1)]
			kept_rows.append((
				[cell.value for cell in cells],
				[copy(cell._style) for cell in cells],
				sheet.row_dimensions[row_index].height,
			))
		if sheet.max_row > 1:
			sheet.delete_rows(2, sheet.max_row - 1)
		for output_index, (values, styles, height) in enumerate(kept_rows, start=2):
			sheet.row_dimensions[output_index].height = height
			for column, (value, style) in enumerate(zip(values, styles), start=1):
				cell = sheet.cell(output_index, column, value)
				if style:
					cell._style = style
		for table in sheet.tables.values():
			table.ref = f'A1:{get_column_letter(len(headers))}{sheet.max_row}'
			if table.autoFilter:
				table.autoFilter.ref = table.ref
		row_count[sheet.title] = sheet.max_row - 1
		if [cell.value for cell in sheet[1]] != original_headers[sheet.title]:
			raise RuntimeError(f'Column headers changed in {sheet.title}')
		if len(headers) != sheet.max_column:
			raise RuntimeError(f'Column count changed in {sheet.title}')
		if row_count[sheet.title] and row_count[sheet.title] < 1:
			raise RuntimeError(f'No selected rows remain in {sheet.title}')
	if len(workbook.sheetnames) != 5:
		raise RuntimeError(f'Expected the original five sheets, got {workbook.sheetnames}')
	if row_count.get('judgments') != sample_size:
		raise RuntimeError(f'Expected {sample_size} judgment rows, got {row_count.get("judgments")}')
	return row_count


def main():
	parser = argparse.ArgumentParser()
	parser.add_argument('--source', type=Path, default=DEFAULT_SOURCE)
	parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
	parser.add_argument('--sample-size', type=int, default=DEFAULT_SAMPLE_SIZE)
	parser.add_argument('--seed', type=int, default=DEFAULT_SEED)
	args = parser.parse_args()
	if args.sample_size < 1:
		raise ValueError('--sample-size must be at least 1')
	if args.source.resolve() == args.output.resolve():
		raise ValueError('Source and output paths must be different')

	workbook = openpyxl.load_workbook(args.source, data_only=False)
	sheets = {sheet.title: read_rows(sheet) for sheet in workbook.worksheets}
	case_rows, eligible, features = build_features(sheets)
	if args.sample_size > len(eligible):
		raise ValueError(f'Sample size {args.sample_size} exceeds eligible cases {len(eligible)}')
	selected = select_cases(features, args.sample_size, args.seed)
	row_counts = filter_workbook(workbook, selected, args.sample_size)
	args.output.parent.mkdir(parents=True, exist_ok=True)
	workbook.save(args.output)
	selected_features = Counter(
		feature
		for case_id in selected
		for feature in features[case_id]
	)
	print(json.dumps({
		'source': str(args.source.resolve()),
		'output': str(args.output.resolve()),
		'eligible_cases': len(eligible),
		'sampled_cases': len(selected),
		'seed': args.seed,
		'sheets': row_counts,
		'age_band_case_counts': {
			band: selected_features['age:' + band]
			for band in ('Under 21', '21–34', '35+')
		},
		'role_case_counts': {
			role: selected_features['role:' + role]
			for role in ROLE_TARGETS
		},
	}, ensure_ascii=False))


if __name__ == '__main__':
	main()
